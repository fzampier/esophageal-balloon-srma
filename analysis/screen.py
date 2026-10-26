"""Title/abstract screening of the search update by two model reviewers.

Same rule as extraction and risk of bias: GLM and DeepSeek get the identical
prompt (studies/search/SCREENING_PROMPT.md) and the identical records, in the
same batches. Where they agree, their decision stands; the PI rules on every
record where they disagree (D40).

    python analysis/screen.py run [--only-new] [--workers 8] [--reviewer glm|deepseek]
    python analysis/screen.py check     # recall on February's 33 included reports, agreement
    python analysis/screen.py list      # records either model kept, as a CSV
    python analysis/screen.py worksheet # disagreements -> xlsx for the PI's ruling
    python analysis/screen.py prisma    # flow counts of the update, from the files (D41, D42)

Answers are cached per reviewer in screening/<reviewer>.jsonl; a re-run only
sends records that have no answer yet.
"""
import argparse
import csv
import json
import sys
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import extract  # noqa: E402  (REVIEWERS, post, parse_record)
import search_update  # noqa: E402
from rob import load_keys  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
DIR = ROOT / "studies" / "search" / "update_2026-09-29"
OUT = DIR / "screening"
PROMPT = ROOT / "studies" / "search" / "SCREENING_PROMPT.md"
BATCH = 10


def records(only_new):
    name = "possibly_new.csv" if only_new else "records_dedup.csv"
    with open(DIR / name, encoding="utf-8") as f:
        return list(csv.DictReader(f))


def record_text(r):
    parts = [f"record_id: {r['record_id']}",
             f"title: {r['title']}",
             f"source: {r['journal']} ({r['year']}); type: {r['pub_type'] or 'not given'}",
             f"abstract: {r['abstract'] or '(no abstract)'}"]
    return "\n".join(parts)


def load_answers(reviewer):
    path = OUT / f"{reviewer}.jsonl"
    if not path.exists():
        return {}
    ans = {}
    for line in path.read_text(encoding="utf-8").splitlines():
        d = json.loads(line)
        ans[d["record_id"]] = d
    return ans


def screen_batch(reviewer, batch, keys):
    cfg = extract.REVIEWERS[reviewer]
    content = (PROMPT.read_text(encoding="utf-8").strip()
               + "\n\n# Records\n\n" + "\n\n---\n\n".join(record_text(r) for r in batch))
    body = {"model": cfg["model"], "messages": [{"role": "user", "content": content}],
            "temperature": 0, "max_tokens": 16384, **cfg["extra"]}
    raw = extract.post(cfg["url"], keys[cfg["key"]], body, timeout=600)
    got = extract.parse_record(raw["choices"][0]["message"].get("content") or "").get("decisions", [])
    want = {r["record_id"] for r in batch}
    out = [d for d in got if d.get("record_id") in want
           and d.get("decision") in ("include", "exclude")]
    for d in out:
        d["reviewer"], d["model"] = reviewer, cfg["model"]
    return out


def cmd_run(args):
    keys = load_keys()
    OUT.mkdir(exist_ok=True)
    recs = records(args.only_new)
    for reviewer in ([args.reviewer] if args.reviewer else extract.PAIRS["both"]):
        done = load_answers(reviewer)
        todo = [r for r in recs if r["record_id"] not in done]
        batches = [todo[i:i + BATCH] for i in range(0, len(todo), BATCH)]
        print(f"{reviewer}: {len(todo)} records to screen in {len(batches)} batches")
        with ThreadPoolExecutor(args.workers) as pool, \
                open(OUT / f"{reviewer}.jsonl", "a", encoding="utf-8") as f:
            futs = {pool.submit(screen_batch, reviewer, b, keys): b for b in batches}
            for n, fut in enumerate(as_completed(futs), 1):
                try:
                    for d in fut.result():
                        f.write(json.dumps(d, ensure_ascii=False) + "\n")
                    f.flush()
                except Exception as exc:  # a failed batch is retried on the next run
                    print(f"  batch failed: {exc}", flush=True)
                if n % 20 == 0:
                    print(f"  {n}/{len(batches)}", flush=True)
    print("run again to fill any gaps, then: screen.py check")


UPDATE_MARK = "Search update 2026-09-29"   # manifest notes of the reports this update added


def included_ids(february_only=False):
    """record_id of each included report in the re-run, matched as in search_update.
    february_only: the 33 reports of the February human screen, the recall reference."""
    kept = records(False)
    by_doi = {r["doi"]: r["record_id"] for r in kept if r["doi"]}
    by_title = {search_update.norm_title(r["title"])[:60]: r["record_id"] for r in kept}
    out = {}
    with open(ROOT / "studies" / "manifest.csv", encoding="utf-8") as f:
        for m in csv.DictReader(f):
            if february_only and m["notes"].startswith(UPDATE_MARK):
                continue
            doi = search_update.norm_doi(m["doi"] or "")
            rid = by_doi.get(doi) if doi.startswith("10.") else None
            out[m["study_id"]] = rid or by_title.get(search_update.norm_title(m["title"])[:60])
    return out


def cmd_check(args):
    a = {r: load_answers(r) for r in extract.PAIRS["both"]}
    g, d = a["glm"], a["deepseek"]
    both = set(g) & set(d)
    print(f"answered: glm {len(g)}, deepseek {len(d)}, both {len(both)}")
    agree = sum(g[i]["decision"] == d[i]["decision"] for i in both)
    inc = lambda x, i: x.get(i, {}).get("decision") == "include"
    kept = [i for i in both if inc(g, i) or inc(d, i)]
    if both:
        po = agree / len(both)
        pg = sum(inc(g, i) for i in both) / len(both)
        pd = sum(inc(d, i) for i in both) / len(both)
        pe = pg * pd + (1 - pg) * (1 - pd)
        kappa = (po - pe) / (1 - pe) if pe < 1 else float("nan")
        print(f"agreement {agree}/{len(both)} ({po:.1%}), Cohen's kappa {kappa:.2f}; "
              f"kept by either {len(kept)}, by both {sum(inc(g, i) and inc(d, i) for i in both)}")
    print("\nrecall on the included reports (February human screen):")
    miss = 0
    for sid, rid in included_ids(february_only=True).items():
        if not rid or rid not in both:
            print(f"  {sid}: not screened ({rid})")
            continue
        gi, di = inc(g, rid), inc(d, rid)
        if not (gi or di):
            miss += 1
        if not (gi and di):
            print(f"  {sid} {rid}: glm {g[rid]['decision']} ({g[rid].get('reason')}), "
                  f"deepseek {d[rid]['decision']} ({d[rid].get('reason')})")
    print(f"included reports excluded by both models: {miss}")


def cmd_list(args):
    a = {r: load_answers(r) for r in extract.PAIRS["both"]}
    new = {r["record_id"] for r in records(True)}
    inc_ids = set(filter(None, included_ids().values()))
    rows = []
    for r in records(False):
        rid = r["record_id"]
        g, d = a["glm"].get(rid, {}), a["deepseek"].get(rid, {})
        if "include" not in (g.get("decision"), d.get("decision")):
            continue
        if args.only_new and rid not in new:
            continue
        rows.append({
            "record_id": rid, "already_included": rid in inc_ids, "possibly_new": rid in new,
            "glm": g.get("decision"), "glm_type": g.get("record_type"), "glm_note": g.get("note"),
            "deepseek": d.get("decision"), "deepseek_type": d.get("record_type"),
            "deepseek_note": d.get("note"),
            **{k: r[k] for k in ("year", "title", "authors", "journal", "doi", "pmid", "nct",
                                 "database", "pub_type", "abstract")}})
    path = OUT / ("kept_new.csv" if args.only_new else "kept_all.csv")
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0]) if rows else ["record_id"])
        w.writeheader()
        w.writerows(rows)
    print(f"{len(rows)} records -> {path.relative_to(ROOT)}")


def cmd_worksheet(args):
    """Every record where the reviewers disagree, for the PI's ruling (D40)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    a = {r: load_answers(r) for r in extract.PAIRS["both"]}
    new = {r["record_id"] for r in records(True)}
    inc_ids = set(filter(None, included_ids().values()))
    wb = Workbook()
    ws = wb.active
    ws.title = "disagreements"
    head = ["record_id", "ruling (include / exclude)", "note", "already_included", "possibly_new",
            "year", "title", "journal", "pub_type", "glm", "glm_note", "deepseek", "deepseek_note",
            "abstract", "doi", "pmid", "nct"]
    ws.append(head)
    fmt = lambda x: x["decision"] + (f" ({x['reason']})" if x.get("reason") else "") + (
        f" [{x['record_type']}]" if x.get("record_type") not in (None, "article") else "")
    n = 0
    for r in records(False):
        rid = r["record_id"]
        g, d = a["glm"].get(rid), a["deepseek"].get(rid)
        if not g or not d or g["decision"] == d["decision"]:
            continue
        ws.append([rid, "", "", "yes" if rid in inc_ids else "", "yes" if rid in new else "",
                   r["year"], r["title"], r["journal"], r["pub_type"], fmt(g), g.get("note"),
                   fmt(d), d.get("note"), r["abstract"], r["doi"], r["pmid"], r["nct"]])
        n += 1
    agreed = wb.create_sheet("agreed includes")
    agreed.append(["record_id", "possibly_new", "year", "title", "journal", "pub_type", "glm",
                   "glm_note", "deepseek", "deepseek_note", "abstract", "doi", "pmid", "nct"])
    for r in records(False):
        rid = r["record_id"]
        g, d = a["glm"].get(rid), a["deepseek"].get(rid)
        if rid in inc_ids or not g or not d or not g["decision"] == d["decision"] == "include":
            continue
        agreed.append([rid, "yes" if rid in new else "", r["year"], r["title"], r["journal"],
                       r["pub_type"], fmt(g), g.get("note"), fmt(d), d.get("note"),
                       r["abstract"], r["doi"], r["pmid"], r["nct"]])
    widths = dict(A=8, B=14, C=24, D=9, E=9, F=6, G=60, H=24, I=18, J=16, K=45, L=16, M=45,
                  N=100, O=22, P=10, Q=13)
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    for c in ws[1]:
        c.font = Font(bold=True)
    for row in ws.iter_rows(min_row=2):
        row[1].fill = PatternFill("solid", fgColor="FFF2CC")
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "C2"
    notes = wb.create_sheet("read me")
    for line in [
        "Search update, 29 September 2026: title/abstract screen by GLM and DeepSeek (D40).",
        "These are the records where the two models disagree. Where they agree, their decision stands.",
        "Ruling: include (goes to full text) or exclude. Exclusion codes: P paediatric/animal, V not invasively",
        "ventilated in ICU, I esophageal pressure not used to set PEEP, O feasibility/technical only,",
        "S fewer than 5 patients, D review/editorial/comment/guideline.",
        "[conference_abstract] and [registry] mark records a model would keep; under D41 both are excluded.",
        "Only possibly-new records are screened (D41); the February human screen stands for the rest.",
        "Sheet 'agreed includes': both models keep it, not already included. Possibly-new articles among",
        "them go to full text without a ruling (D40, D41).",
    ]:
        notes.append([line])
    path = OUT / args.out
    wb.save(path)
    print(f"{n} disagreements -> {path.relative_to(ROOT)}")


def cmd_prisma(args):
    """PRISMA counts of the update, computed from the answers and the PI's rulings."""
    raw = sum(len(search_update.parse_ris(DIR / n)) for n in search_update.FILES.values())
    alln, new = records(False), records(True)
    a = {r: load_answers(r) for r in extract.PAIRS["both"]}
    feb = set(filter(None, included_ids(february_only=True).values()))
    pi = {r["record_id"]: r for r in csv.DictReader(open(OUT / "pi_rulings.csv", encoding="utf-8"))}
    ft = {r["record_id"]: r for r in csv.DictReader(open(OUT / "fulltext_decisions.csv", encoding="utf-8"))}
    c, to_full = {}, []
    bump = lambda k: c.__setitem__(k, c.get(k, 0) + 1)
    for r in new:
        rid = r["record_id"]
        if rid in feb:
            bump("already included in February")
            continue
        g, d = a["glm"][rid], a["deepseek"][rid]
        types = {g.get("record_type"), d.get("record_type")}
        if rid in pi:
            (to_full.append(rid) if pi[rid]["ruling"] == "include" else bump("excluded: models disagree, PI excludes"))
        elif g["decision"] == d["decision"] == "exclude":
            bump("excluded: both models exclude")
        elif "registry" in types:
            bump("excluded: trial registry entry (D41)")
        elif "conference_abstract" in types:
            bump("excluded: conference abstract (D41)")
        elif g["decision"] == d["decision"] == "include":
            to_full.append(rid)
        else:
            bump("UNRULED DISAGREEMENT")
    screened = len(new) - c.get("already included in February", 0)
    excl = sum(v for k, v in c.items() if k.startswith("excluded"))
    print(f"records identified {raw}; duplicates removed {raw - len(alln)}; unique {len(alln)}")
    print(f"screened in February, not re-screened {len(alln) - len(new)}; possibly new {len(new)}")
    for k, v in sorted(c.items()):
        print(f"  {k}: {v}")
    print(f"records screened {screened}; excluded at title/abstract {excl}; full texts {len(to_full)}")
    missing = [i for i in to_full if i not in ft]
    inc = [i for i in to_full if ft.get(i, {}).get("decision") == "include"]
    print(f"full text included {len(inc)}; excluded {len(to_full) - len(inc) - len(missing)}; no decision {missing}")
    # counts for the PRISMA figure (docs/manuscript/make_figures.R): the model screen, shown until the
    # authors' screen of record (D51) sends its own counts (D56)
    reason = {"D": "Perspective article", "I": "Esophageal pressure did not set PEEP",
              "X": "Secondary analysis without new patients", "O": "No patient outcomes"}
    excl_ft = [ft[i]["reason"] for i in to_full if ft.get(i, {}).get("decision") == "exclude"]
    reasons = "; ".join(f"{reason[k]}, {excl_ft.count(k)}" for k in reason if excl_ft.count(k))
    rows = [("database_results", raw), ("duplicates", raw - len(alln)), ("excluded_other", len(alln) - screened),
            ("records_screened", screened), ("records_excluded", excl), ("dbr_sought_reports", len(to_full)),
            ("dbr_notretrieved_reports", 0), ("dbr_assessed", len(to_full)), ("dbr_excluded", reasons),
            ("new_reports", len(inc))]
    with open(DIR / "prisma_counts_model_screen.csv", "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["data", "n"])
        w.writerows(rows)
    print(f"wrote {(DIR / 'prisma_counts_model_screen.csv').relative_to(ROOT)}")


def main():
    ap = argparse.ArgumentParser()
    sub = ap.add_subparsers(dest="cmd", required=True)
    p = sub.add_parser("run")
    p.add_argument("--only-new", action="store_true")
    p.add_argument("--workers", type=int, default=8)
    p.add_argument("--reviewer", choices=extract.PAIRS["both"])
    sub.add_parser("check")
    p = sub.add_parser("worksheet")
    p.add_argument("--out", default="screening_disagreements.xlsx")
    sub.add_parser("prisma")
    p = sub.add_parser("list")
    p.add_argument("--only-new", action="store_true")
    args = ap.parse_args()
    {"run": cmd_run, "check": cmd_check, "list": cmd_list, "worksheet": cmd_worksheet, "prisma": cmd_prisma}[args.cmd](args)


if __name__ == "__main__":
    main()

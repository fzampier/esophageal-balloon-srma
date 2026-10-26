#!/usr/bin/env python3
"""Build a spot-check workbook: a person checks reconciled values against the PDFs.

    python analysis/spotcheck.py --checker "Imran Hudda" sun_2024 wang_2020 ...

One row per value that reaches the analysis (analysis/scope.py's footprint, as
flatten.py and analysis/meta/ read it): the study's size and subgroup classes,
every mortality entry, ventilator-free days, durations, stays, barotrauma, and
the physiology measures the D14 tables show. Each row carries the record's own
quote and page location, so the checker opens the PDF at that spot and answers
Correct / Wrong / Can't find. Nothing here writes to a record; a Wrong comes
back as a proposed post-merge edit (C17) for the PI to rule on.

    python analysis/spotcheck.py --scope core --all --checker "Anton Nikouline" \
        --out extraction/manual/check/Esophageal_balloon_SRMA_check_Anton_2026-10-23.xlsx

`--scope core` is the authors' check of the key values (D85): every analysed
report, and only its design, the numbers enrolled, analysed and per group, and
the six clinical outcomes; no baseline, subgroup classes or physiology. Two
authors each fill an identical workbook, independently.
"""

import argparse
import csv
import json
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from manual_form import LISTS

ROOT = Path(__file__).resolve().parent.parent
REC = ROOT / "extraction" / "reconciled"
OUT = ROOT / "extraction" / "adjudication" / "spotcheck"

CLINICAL = {
    "mortality": "Mortality",
    "ventilator_free_days": "Ventilator-free days",
    "duration_mechanical_ventilation": "Duration of mechanical ventilation",
    "icu_length_of_stay": "ICU length of stay",
    "hospital_length_of_stay": "Hospital length of stay",
    "barotrauma": "Barotrauma",
}
MEASURE = {
    "peep_set": "PEEP applied", "peep_total": "Total PEEP", "pao2_fio2": "PaO2/FiO2",
    "oxygenation_index": "Oxygenation index", "compliance_rs": "Respiratory-system compliance",
    "elastance_rs": "Respiratory-system elastance", "plateau_pressure": "Plateau pressure",
    "driving_pressure": "Driving pressure", "pl_end_expiratory": "End-expiratory transpulmonary pressure",
    "pl_driving": "Transpulmonary driving pressure",
}
ROLE = {"intervention": "balloon-guided group", "comparator": "comparison group",
        "reference_condition": "reference (conventional) condition", "all_patients": "all patients"}
SUBGROUP = {"ards_subgroup": "ARDS population (subgroup class)",
            "obesity_subgroup": "Obesity (subgroup class)",
            "timing_subgroup": "Timing of balloon use (subgroup class)",
            "ptp_target_subgroups": "Transpulmonary pressure target (subgroup class)"}

COLUMNS = ["#", "study", "PDF file", "section", "what this is", "group", "timepoint",
           "value we recorded", "where we read it", "the paper's words we read it from",
           "CHECK", "if wrong: what the paper says (and page)", "comment"]
ANSWER = PatternFill("solid", fgColor="FFF2CC")
BOLD = Font(bold=True)


BASELINE = {
    "age": "Age",
    "sex_female": "Women (number / group size)",
    "bmi": "Body mass index",
    "pao2_fio2_baseline": "PaO2/FiO2 at baseline",
    "apache_ii": "APACHE II",
    "sofa": "SOFA",
}


def measure_map():
    with (ROOT / "analysis" / "meta" / "measure_map.csv").open(encoding="utf-8") as fh:
        return {(r["study_id"], r["name_verbatim"]): r["measure"] for r in csv.DictReader(fh)
                if r["measure"] != "not_tabulated"}


def pdfs():
    with (ROOT / "studies" / "manifest.csv").open(encoding="utf-8") as fh:
        return {r["study_id"]: r["filename"] for r in csv.DictReader(fh)}


def num(x):
    return "" if x is None else (f"{x:g}" if isinstance(x, float) else str(x))


def show_value(d):
    if d.get("events") is not None:
        return f"{num(d['events'])} / {num(d.get('total'))}"
    if d.get("value") is None:
        n = d.get("total") if d.get("total") is not None else d.get("n")
        return "no value recorded" + (f" (n = {num(n)})" if n is not None else "")
    parts = [f"{d.get('value_type') or 'value'} {num(d.get('value'))}"]
    dt = d.get("dispersion_type")
    if d.get("dispersion") is not None:
        parts.append(f"({dt} {num(d['dispersion'])})")
    elif d.get("dispersion_low") is not None or d.get("dispersion_high") is not None:
        parts.append(f"({dt} {num(d.get('dispersion_low'))} to {num(d.get('dispersion_high'))})")
    if d.get("unit"):
        parts.append(d["unit"])
    if d.get("n") is not None:
        parts.append(f"n = {num(d['n'])}")
    return " ".join(parts)


def timepoint(o):
    t = o.get("timepoint") or {}
    if t.get("label_verbatim"):
        return t["label_verbatim"]
    if t.get("days") is not None:
        return f"day {num(t['days'])}"
    if t.get("hours") is not None:
        return f"{num(t['hours'])} h"
    return (t.get("type") or "").replace("_", " ")


DESIGN = {code: label for label, code in LISTS["design"]}
STRUCTURE = {code: label for label, code in LISTS["structure"]}


def analysed():
    with (ROOT / "analysis" / "derived" / "studies.csv").open(encoding="utf-8") as fh:
        return sorted(r["study_id"] for r in csv.DictReader(fh) if r["analysed"] == "True")


def rows_for(sid, mm, core=False):
    rec = json.loads((REC / f"{sid}.json").read_text(encoding="utf-8"))
    arms = {a["label"]: a for a in rec.get("arms", [])}

    def group(label):
        a = arms.get(label, {})
        words = a.get("label_verbatim") or label
        role = ROLE.get(a.get("role"), a.get("role") or "")
        return f"{words} ({role})" if role else words

    out = []
    if core:
        des = rec.get("design") or {}
        p = (des.get("study_design_verbatim") or {}).get("provenance") or {}
        code = des.get("study_design")
        out.append(("Study", "Study design", "whole study", "", DESIGN.get(code, code or ""),
                    p.get("location"), p.get("quote")))
        code = des.get("comparison_structure")
        out.append(("Study", "Comparison", "whole study", "", STRUCTURE.get(code, code or ""),
                    "our classification", None))
    pop = rec.get("population", {})
    for key, what in (("n_enrolled", "Patients enrolled"), ("n_analysed", "Patients analysed")):
        v = pop.get(key)
        if isinstance(v, dict) and v.get("value") is not None:
            p = v.get("provenance") or {}
            out.append(("Study", what, "whole study", "", num(v["value"]), p.get("location"), p.get("quote")))
    # in a within-patient study the "arms" are conditions of the same patients:
    # their sizes are not group sizes, so the core check leaves them out (D85)
    parallel = (rec.get("design") or {}).get("comparison_structure") == "parallel_arms"
    for a in rec.get("arms", []) if parallel or not core else []:
        for key, what in (("n_randomised", "Patients in group (randomised/allocated)"),
                          ("n_analysed", "Patients in group (analysed)")):
            v = a.get(key)
            if isinstance(v, dict) and v.get("value") is not None:
                p = v.get("provenance") or {}
                out.append(("Study", what, group(a["label"]), "", num(v["value"]), p.get("location"), p.get("quote")))
    # Baseline characteristics (B8, C25): whole study where the paper gives an
    # overall figure, by group where Table 1 is by group.
    def baseline(node, who):
        for key, what in BASELINE.items():
            v = node.get(key)
            if not isinstance(v, dict):
                continue
            p = v.get("provenance") or {}
            if key == "sex_female":
                if v.get("n") is not None:
                    out.append(("Baseline", what, who, "", f"{num(v['n'])} / {num(v.get('total'))}",
                                p.get("location"), p.get("quote")))
            elif v.get("value") is not None:
                out.append(("Baseline", what, who, "", show_value(v), p.get("location"), p.get("quote")))

    if not core:
        baseline(pop, "whole study")
        for arm in pop.get("baseline_by_arm") or []:
            baseline(arm, group(arm.get("arm_label")))

    flags = {} if core else rec.get("subgroup_flags", {})
    for key, what in SUBGROUP.items():
        v = flags.get(key)
        if v not in (None, [], ""):
            v = "; ".join(v) if isinstance(v, list) else v
            out.append(("Study", what, "whole study", "", v.replace("_", " "), "our classification",
                        flags.get("classification_notes")))

    for o in rec.get("outcomes", []):
        dom, name = o.get("domain"), o.get("name_verbatim")
        if dom in CLINICAL:
            section = "Clinical outcome"
            what = f"{CLINICAL[dom]}: {name}"
            if dom == "mortality" and o.get("is_longest_reported_mortality"):
                what += "  [USED AS THE MAIN MORTALITY RESULT]"
        elif (sid, name) in mm and not core:
            section = "Physiology"
            what = f"{MEASURE.get(mm[(sid, name)], mm[(sid, name)])}: {name}"
        else:
            continue
        for d in o.get("data", []):
            p = d.get("provenance") or {}
            out.append((section, what, group(d.get("arm_label")), timepoint(o), show_value(d),
                        p.get("location"), p.get("quote")))
    return rec, out


CORE_TEXT = [
    "What this is. Two AI readers extracted every paper, and their disagreements were settled by the "
    "authors and a third model. Before submission, two of us each check, on our own, every value the "
    "pooled analyses use against the papers themselves: the clinical outcomes (mortality, duration of "
    "ventilation, ventilator-free days, ICU and hospital stay, barotrauma), the study design, and the "
    "number of patients in the study and in each group. Baseline characteristics and physiology are "
    "not part of this check.",
    "",
    "Please work alone and do not compare answers until both of you have finished; Fernando goes "
    "through any differences with you afterwards. Get the papers (and supplements) through your own "
    "library access: the DOI of each is in the sheet 'Studies'.",
    "",
    "What to do.",
    "1. Open the sheet 'Check'. Rows are grouped by study.",
    "2. For each row, find the spot named in 'where we read it' and compare 'value we recorded' "
    "with the paper. 'The paper's words' is the text the reader copied, to help you find it.",
    "3. In CHECK choose: Correct / Wrong / Can't find. If Wrong, write what the paper says and the "
    "page or table in the next column. Anything else goes in 'comment'.",
    "",
    "What matters most. The rows marked [USED AS THE MAIN MORTALITY RESULT], the group sizes, and "
    "which group is the balloon-guided one. A wrong group label (numbers under the wrong arm) "
    "is the error we most want to catch.",
    "Also check: is the mortality timepoint the longest follow-up the paper reports? Is a clinical "
    "result missing (a mortality, ventilator-free-days, duration, length-of-stay or barotrauma result "
    "that is in the paper but not in the sheet)? Note it in 'comment' on any row of that study.",
    "",
    "Design: is it the design the paper describes (a randomized trial, a cohort, a physiologic study in "
    "which each patient had both strategies)? 'Comparison' says whether the groups are different patients "
    "or the same patients under each strategy.",
    "",
    "Values are what the paper reports: nothing converted, no medians turned into means. "
    "'iqr 0 to 47' means the interquartile range the paper prints.",
    "",
    "When done, save the file and send it back to Fernando.",
]


def build(study_ids, checker, path, core=False):
    mm, files = measure_map(), pdfs()
    wb = Workbook()
    ws = wb.active
    ws.title = "How to check"
    text = [f"Check of the key extracted values — for {checker}", "", *CORE_TEXT] if core else [
        f"Spot check of the extracted data — for {checker}",
        "",
        "What this is. Two AI readers extracted every paper; their disagreements were settled by the "
        "co-authors and a third model. Before we write, we want a person to check a sample of the "
        "final numbers against the papers themselves.",
        "",
        "What to do.",
        "1. Open the sheet 'Check'. Rows are grouped by study; the PDF file name is on every row "
        "(Fernando can send any PDF you do not have).",
        "2. For each row, find the spot named in 'where we read it' and compare 'value we recorded' "
        "with the paper. 'The paper's words' is the text the reader copied, to help you find it.",
        "3. In CHECK choose: Correct / Wrong / Can't find. If Wrong, write what the paper says and the "
        "page or table in the next column. Anything else goes in 'comment'.",
        "",
        "What matters most. The rows marked [USED AS THE MAIN MORTALITY RESULT], the group sizes, and "
        "which group is the balloon-guided one. A wrong group label (numbers under the wrong arm) "
        "is the error we most want to catch.",
        "Also useful: is the mortality timepoint the longest follow-up the paper reports? Is anything "
        "clinically important missing (a mortality, ventilator-free-days, length-of-stay or barotrauma "
        "result that is in the paper but not in the sheet)? Note it in 'comment' on any row of that study.",
        "",
        "Subgroup rows show how we classified the study (ARDS, obesity, timing, pressure target), with "
        "our reasoning in the last column. Tell us if you disagree.",
        "",
        "Values are what the paper reports: nothing converted, no medians turned into means. "
        "'iqr 0 to 47' means the interquartile range the paper prints.",
        "",
        "When done, save the file with your name and send it back to Fernando.",
    ]
    for i, line in enumerate(text, 1):
        c = ws.cell(row=i, column=1, value=line)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if i == 1:
            c.font = Font(bold=True, size=13)
    ws.column_dimensions["A"].width = 120

    st = wb.create_sheet("Studies")
    st.append(["study", "first author", "year", "journal", "title", "design", "PDF file", "rows to check", "DOI"])
    ck = wb.create_sheet("Check")
    ck.append(COLUMNS)
    n = 0
    for sid in study_ids:
        rec, rows = rows_for(sid, mm, core)
        c = rec["citation"]
        st.append([sid, c.get("first_author"), c.get("year"), c.get("journal"), c.get("title"),
                   (rec.get("design") or {}).get("study_design"), files.get(sid), len(rows), c.get("doi")])
        if c.get("doi"):
            st.cell(row=st.max_row, column=9).hyperlink = f"https://doi.org/{c['doi']}"
        for r in rows:
            n += 1
            ck.append([n, sid, files.get(sid), *r, "", "", ""])

    for sheet in (st, ck):
        for cell in sheet[1]:
            cell.font = BOLD
        sheet.freeze_panes = "A2"
    widths = [5, 14, 22, 14, 42, 30, 18, 26, 30, 50, 12, 36, 30]
    for i, w in enumerate(widths, 1):
        ck.column_dimensions[get_column_letter(i)].width = w
    for row in ck.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for cell in row[10:]:
            cell.fill = ANSWER
    dv = DataValidation(type="list", formula1='"Correct,Wrong,Can\'t find"', allow_blank=True)
    ck.add_data_validation(dv)
    dv.add(f"K2:K{n + 1}")
    ck.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{n + 1}"
    for i, w in enumerate([14, 16, 6, 24, 60, 20, 30, 12, 30], 1):
        st.column_dimensions[get_column_letter(i)].width = w

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return n


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("study_ids", nargs="*")
    ap.add_argument("--all", action="store_true", help="every analysed report (A4)")
    ap.add_argument("--scope", choices=["sample", "core"], default="sample",
                    help="core: design, group sizes and clinical outcomes only (D85)")
    ap.add_argument("--checker", required=True)
    ap.add_argument("--out", type=Path)
    args = ap.parse_args()
    ids = analysed() if args.all else args.study_ids
    if not ids:
        ap.error("name the studies or pass --all")
    out = args.out or OUT / f"spotcheck_{args.checker.split()[0]}_{date.today().isoformat()}.xlsx"
    n = build(ids, args.checker, out, args.scope == "core")
    print(f"{out}: {n} row(s) across {len(ids)} studies")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

#!/usr/bin/env python3
"""The adjudication ledger: every ruling on a reviewer difference, and who made it.

    python3 analysis/rulings.py import returned.xlsx --adjudicator "Anton Nikouline"
    python3 analysis/rulings.py check               # rulings that no longer match the diff

A returned workbook is not kept as the record of adjudication — it is gitignored,
arrives in several copies, and its rows are keyed by position. Importing it turns
each DECISION into a line of extraction/adjudication/rulings.csv (field-level) or
study_rulings.csv (mortality flag, cohort relations), keyed by what the ruling is
about: study, field, kind, and both reviewers' readings as diff.py shows them.
reconcile.py looks rulings up by that key, so a ruling follows its difference
across workbook rebuilds and is reported as orphaned, never silently misapplied,
if the difference it settled stops existing.

Workbooks built before diff.py compared timepoints (commit fd318a7) name fields
without the `@72h` suffix. Those rows are matched with the suffix stripped, and
the ruling is recorded against every current row that matches — which is only
sound because the readings must match too.
"""

import argparse
import csv
import hashlib
import re
import sys
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import worksheet as W  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
LEDGER = ROOT / "extraction" / "adjudication"
RULINGS = LEDGER / "rulings.csv"
STUDY_RULINGS = LEDGER / "study_rulings.csv"
COHORTS = LEDGER / "cohorts.csv"
COHORT_COLUMNS = ["study_id", "status", "cohort_id", "overlapping_cohort_with",
                  "is_primary_report", "basis", "adjudicator", "note"]

RULING_COLUMNS = [
    "ruling_id", "study_id", "field", "kind", "glm", "deepseek",
    "decision", "value", "note", "adjudicator", "adjudicator_type", "basis",
    "source", "source_field", "recorded",
]
STUDY_COLUMNS = [
    "ruling_id", "study_id", "topic", "decision", "value", "note",
    "adjudicator", "adjudicator_type", "basis", "source", "recorded",
]

# DECISION vocabulary on the workbook -> ledger vocabulary.
DECISIONS = {
    "glm": "glm",
    "deepseek": "deepseek",
    "both wrong": "value",
    "not reported": "not_reported",
    "unclear": "unclear",
}


def strip_timepoints(field):
    return re.sub(r" @[^\]\[]*\]", "]", field)


def ruling_id(*parts):
    return hashlib.sha1("|".join(str(p) for p in parts).encode()).hexdigest()[:12]


def read_csv(path):
    if not path.exists():
        return []
    with path.open(newline="", encoding="utf-8") as fh:
        return list(csv.DictReader(fh))


def write_csv(path, columns, rows):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=columns, lineterminator="\n")
        w.writeheader()
        w.writerows(rows)


def load(ledger=LEDGER):
    """(field rulings, study rulings), in the order they were recorded."""
    return read_csv(ledger / RULINGS.name), read_csv(ledger / STUDY_RULINGS.name)


def row_key(study_id, field, kind, glm, deepseek):
    return (study_id, field, kind, str(glm), str(deepseek))


def sheet_rows(ws):
    rows = list(ws.iter_rows(values_only=True))
    head = [str(h) if h is not None else "" for h in rows[0]]
    for i, r in enumerate(rows[1:], start=2):
        if any(c not in (None, "") for c in r):
            yield i, {h: ("" if c is None else c) for h, c in zip(head, r)}


def cell(v):
    return "" if v is None else str(v).strip()


def import_workbook(path, adjudicator, adjudicator_type="human"):
    from openpyxl import load_workbook

    digest = hashlib.sha256(Path(path).read_bytes()).hexdigest()[:16]
    rulings, study_rulings = load()
    if any(digest in r["source"] for r in rulings + study_rulings):
        print(f"{path} (sha256 {digest}) is already in the ledger; nothing imported.")
        return 0

    wb = load_workbook(path)
    stamp = cell(wb["Start here"]["A3"].value)
    source = f"{Path(path).name} sha256:{digest} [{stamp}]"
    today = date.today().isoformat()

    current = W.triaged_rows(W.load_manifest())
    exact, legacy = {}, {}
    for r in current:
        k = (r["study_id"], r["field"], r["kind"], r["glm"], r["deepseek"])
        exact.setdefault(k, []).append(r)
        legacy.setdefault(
            (r["study_id"], strip_timepoints(r["field"]), r["kind"], r["glm"], r["deepseek"]), []
        ).append(r)

    # The first workbook's sheet 3 holds only conflicts; the follow-up workbook's
    # sheet 2 names each row's kind, including whole elements the model left unclear.
    data_sheet = next(s for s in ("3. Data conflicts", "2. Readings to rule") if s in wb.sheetnames)
    fields_now = {(r["study_id"], r["field"]) for r in current}
    new, orphans, unclear, clashes, manual = [], [], [], [], []
    for n, row in sheet_rows(wb[data_sheet]):
        decision = DECISIONS.get(cell(row["DECISION"]).lower())
        if not decision:
            continue
        sid, field = cell(row["study_id"]), cell(row["field"])
        glm, ds = cell(row["glm"]), cell(row["deepseek"])
        kind = cell(row.get("kind")) or "conflict"
        if kind not in ("conflict", "one_sided", "element"):
            # A comment on an earlier ruling or a consistency check: the answer
            # is applied by hand as a new ruling on the right difference.
            manual.append((n, sid, field, decision, cell(row["Adjudicated value"]), cell(row["Adjudicator notes"])))
            continue
        if kind == "element":
            live = any(s == sid and (f == field or f.startswith((field + ".", field + "[")))
                       for s, f in fields_now)
            if decision == "value" or not live:
                manual.append((n, sid, field, decision, cell(row["Adjudicated value"]), cell(row["Adjudicator notes"])))
                continue
            targets, glm, ds = [{"field": field}], "", ""
        else:
            targets = exact.get((sid, field, kind, glm, ds)) or legacy.get(
                (sid, strip_timepoints(field), kind, glm, ds), []
            )
        value = cell(row["Adjudicated value"])
        note = cell(row["Adjudicator notes"])
        if decision == "value" and not value:
            # 'Both wrong' with the value column empty settles nothing; it is
            # recorded as Unclear so the row comes back with its note.
            decision, note = "unclear", f"{note} [Both wrong, no value given]".strip()
        chosen = {"glm": glm, "deepseek": ds}.get(decision)
        if chosen is not None and value and value.lower() != chosen.lower():
            # A value typed beside a GLM/DeepSeek ruling restates or summarises
            # the chosen reading; the ruling is the choice, the text is kept.
            note = f"{note} [typed alongside the ruling: {value}]".strip()
        if decision != "value":
            value = ""
        if not targets:
            orphans.append((n, sid, field, decision))
            continue
        if decision == "unclear":
            unclear.append((sid, field))
        for t in targets:
            rid = ruling_id(sid, t["field"], glm, ds, adjudicator, digest)
            same = next((x for x in new if x["ruling_id"] == rid), None)
            if same and (same["decision"], same["value"]) == (decision, value):
                # The old workbook listed one difference twice (the same field at
                # two timepoints, before timepoints were shown); one ruling covers it.
                same["source"] += f", row {n}"
                continue
            if same:
                clashes.append((n, sid, t["field"]))
            new.append({
                "ruling_id": rid if not same else f"{rid}-{n}",
                "study_id": sid, "field": t["field"], "kind": kind,
                "glm": glm, "deepseek": ds, "decision": decision, "value": value,
                "note": note, "adjudicator": adjudicator,
                "adjudicator_type": adjudicator_type, "basis": f"workbook {data_sheet}",
                "source": f"{source} row {n}", "source_field": field, "recorded": today,
            })

    new_study = []
    for sheet, topic in (("1. Mortality flag", "mortality_flag"), ("2. Cohort relations", "cohort")):
        if sheet not in wb.sheetnames:
            continue
        for n, row in sheet_rows(wb[sheet]):
            decision = DECISIONS.get(cell(row.get("DECISION")).lower())
            if not decision:
                continue
            sid = cell(row["study_id"])
            new_study.append({
                "ruling_id": ruling_id(sid, topic, adjudicator, digest),
                "study_id": sid, "topic": topic, "decision": decision,
                "value": cell(row.get("Adjudicated value")),
                "note": cell(row.get("Adjudicator notes")),
                "adjudicator": adjudicator, "adjudicator_type": adjudicator_type,
                "basis": f"workbook {sheet}", "source": f"{source} row {n}", "recorded": today,
            })

    settled_cohorts = []
    if "1. Cohort assignment" in wb.sheetnames:
        cohorts = {r["study_id"]: r for r in read_csv(COHORTS)}
        for n, row in sheet_rows(wb["1. Cohort assignment"]):
            if cell(row.get("DECISION")).lower() != "settled":
                continue
            sid = cell(row["study_id"])
            cohorts[sid] = {
                "study_id": sid, "status": "settled", "cohort_id": cell(row.get("cohort_id")),
                "overlapping_cohort_with": ";".join(
                    x.strip() for x in re.split(r"[;,]", cell(row.get("overlapping_cohort_with"))) if x.strip()),
                "is_primary_report": "true" if cell(row.get("is_primary_report")).lower() in ("true", "yes", "1") else "false",
                "basis": f"{source} row {n}", "adjudicator": adjudicator,
                "note": cell(row.get("Adjudicator notes")),
            }
            settled_cohorts.append(sid)
        write_csv(COHORTS, COHORT_COLUMNS, list(cohorts.values()))

    write_csv(RULINGS, RULING_COLUMNS, rulings + new)
    write_csv(STUDY_RULINGS, STUDY_COLUMNS, study_rulings + new_study)
    print(f"imported {path}")
    print(f"  source   {source}")
    print(f"  field rulings    {len(new)} ledger line(s)  "
          f"({sum(1 for r in new if r['decision'] != 'unclear')} decisive, "
          f"{len(unclear)} returned as Unclear)")
    print(f"  study rulings    {len(new_study)}")
    print(f"  orphaned         {len(orphans)} — the difference no longer exists in the current diff")
    for n, sid, field, decision in orphans:
        print(f"    row {n:<4} {sid:<20} {decision:<9} {field}")
    print(f"  cohorts settled  {len(settled_cohorts)} {', '.join(settled_cohorts)}")
    print(f"  apply by hand    {len(manual)} — comments on earlier rulings, or element rows ruled 'Both wrong'")
    for n, sid, field, decision, value, note in manual:
        print(f"    row {n:<4} {sid:<20} {decision:<9} {field} :: {value} {note}")
    print(f"  clashing         {len(clashes)} — two rows rule the same difference differently")
    for n, sid, field in clashes:
        print(f"    row {n:<4} {sid:<20} {field}")
    return 0


def check():
    rulings, _ = load()
    current = W.triaged_rows(W.load_manifest())
    live = {row_key(r["study_id"], r["field"], r["kind"], r["glm"], r["deepseek"]) for r in current}
    fields = {(r["study_id"], r["field"]) for r in current}

    def alive(r):
        if r["kind"] == "element":  # live while any difference lies beneath the element
            return any(sid == r["study_id"] and (f == r["field"] or f.startswith((r["field"] + ".", r["field"] + "[")))
                       for sid, f in fields)
        return row_key(r["study_id"], r["field"], r["kind"], r["glm"], r["deepseek"]) in live

    stale = [r for r in rulings if not alive(r)]
    print(f"{len(rulings)} field ruling(s); {len(stale)} no longer match a current difference")
    for r in stale:
        print(f"  {r['study_id']:<20} {r['decision']:<9} {r['field']}")
    return 1 if stale else 0


def main():
    ap = argparse.ArgumentParser()
    sub = ap.add_subparsers(dest="cmd", required=True)
    imp = sub.add_parser("import", help="record a returned workbook's DECISIONs")
    imp.add_argument("workbook")
    imp.add_argument("--adjudicator", required=True, help="who made these rulings")
    imp.add_argument("--adjudicator-type", default="human", choices=["human", "model"])
    sub.add_parser("check", help="list rulings whose difference no longer exists")
    args = ap.parse_args()
    if args.cmd == "import":
        return import_workbook(args.workbook, args.adjudicator, args.adjudicator_type)
    return check()


if __name__ == "__main__":
    sys.exit(main())

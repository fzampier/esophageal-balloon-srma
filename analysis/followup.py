#!/usr/bin/env python3
"""Build the follow-up workbook: everything in scope that only a co-author can settle.

    python3 analysis/followup.py                       # -> adjudication_followup.xlsx
    python3 analysis/followup.py --out "~/OneDrive/Desktop/for_Anton.xlsx"

The reader is a clinician co-author who has not read the schema. The first
version of this workbook forgot that: its cohort "questions" were the ledger's
own audit notes, and rows a co-author had returned Unclear — usually to ask what
a field meant — went back to them unanswered, marked "returned Unclear by a
co-author". It was not understood. Every row here says in plain words what it is
about and what is needed. The schema paths and the reviewers' raw readings stay
in the right-hand columns because `rulings.py import` keys on them.

Inputs beside the ledger, all under extraction/adjudication/followup/:

  cohort_briefs.json     per report of a pending cohort: what the paper says
                         (registration, enrolment window, centre, population, the
                         sentence on shared patients) and a suggested assignment
                         for the co-author to ratify or change
  clarifications.json    the answer to each question a co-author asked by
                         returning a row Unclear; the row goes back with it
  concerns.json          extra sheet-2 rows: earlier rulings the model pass found
                         inconsistent with the paper, and consistency checks
  questions.json         class-wide coding questions, in plain words

Sheets, and how each comes back through `rulings.py import`:

  1. Cohort assignment   one row per report of a pending cohort; import reads
                         cohort_id, overlapping_cohort_with, is_primary_report
                         where DECISION = Settled
  2. Readings to rule    import reads study_id, field, kind, glm, deepseek,
                         DECISION, Adjudicated value, Adjudicator notes
  3. Questions           answered in words, applied by hand
  4. Your earlier Unclear rows   information only: what became of each row the
                         co-author returned Unclear in the first workbook
"""

import argparse
import json
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).resolve().parent))
import diff as D  # noqa: E402
import model_pass as MP  # noqa: E402
import reconcile as RC  # noqa: E402
import rulings as R  # noqa: E402
import worksheet as W  # noqa: E402
from scope import in_scope  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
FOLLOWUP = R.LEDGER / "followup"
RETURNED = R.LEDGER / "returned"

SUGGEST = PatternFill("solid", fgColor="EDEDED")
MACHINE = Font(color="808080", size=8)
BOLD = Font(bold=True)

# Header names rulings.py import reads. Never rename these.
READING_COLUMNS = ["study_id", "author", "what this row is about", "what we need from you",
                   "GLM says", "DeepSeek says", "glm_quote", "deepseek_quote"]
READING_ANSWERS = ["DECISION", "Adjudicated value", "Adjudicator notes"]
READING_MACHINE = ["kind", "field", "why_here", "glm", "deepseek"]
COHORT_COLUMNS = ["group", "study_id", "author", "paper", "registration", "enrolment window",
                  "centre", "population", "what the paper says about shared patients",
                  "your earlier answer", "question for you", "our suggestion"]
COHORT_ANSWERS = ["cohort_id", "overlapping_cohort_with", "is_primary_report", "DECISION",
                  "Adjudicator notes"]
EARLIER_COLUMNS = ["study_id", "author", "what the row was about", "GLM said", "DeepSeek said",
                   "your note", "what happened to it"]

HOW_TO_ANSWER = ("If one version is right after all, choose that reader. If the note is right and "
                 "neither is, choose 'Both wrong' and write the fix in 'Adjudicated value' — one line "
                 "is enough, e.g. \"DeepSeek's row, but timepoint = at protocol start\".")


# ---------------------------------------------------------------- plain words

LEAF = {
    "value": "the number itself",
    "dispersion": "the spread (SD, IQR or similar)",
    "dispersion_low": "the lower end of the spread",
    "dispersion_high": "the upper end of the spread",
    "dispersion_type": "what kind of spread it is (SD, IQR, ...)",
    "value_type": "what kind of number it is (mean, median, ...)",
    "events": "the number of patients with the event",
    "total": "the number of patients at risk (the denominator)",
    "n": "the number of patients measured",
    "unit": "the unit",
    "derived": "whether we computed it rather than read it",
    "timepoint.hours": "when it was measured, in hours after the start",
    "timepoint.days": "when it was measured, in days",
    "timepoint.type": "what kind of timepoint this is",
    "timepoint.label_verbatim": "how the paper names the timepoint",
    "is_primary_outcome": "the flag 'this is the paper's primary outcome'",
    "is_longest_reported_mortality": "the flag 'this is the latest mortality the paper reports'",
    "domain": "which category the outcome belongs to",
    "name_verbatim": "the outcome's name as printed",
    "definition_verbatim": "the outcome's definition as printed",
    "measure_type": "continuous or yes/no outcome",
    "arm_label": "which arm or condition the number belongs to",
    "role": "the arm's role (intervention, comparator, condition)",
    "label": "the arm's short name",
    "n_analysed": "the number of patients analysed in the arm",
    "n_randomised": "the number of patients randomised to the arm",
    "sequence_position": "where the condition came in the order applied",
    "p_value": "the P value",
    "p_value_verbatim": "the P value as printed",
    "mean_difference": "the difference between the two conditions",
    "contrast": "which two conditions are compared",
    "contrast_type": "the kind of comparison",
    "arm_label_a": "the first condition of the comparison",
    "arm_label_b": "the second condition of the comparison",
    "test_used": "the statistical test",
    "type": "the kind of effect estimate",
    "ci_low": "the lower confidence limit",
    "ci_high": "the upper confidence limit",
    "adjusted": "whether the estimate is adjusted",
    "adjustment_covariates": "what it was adjusted for",
    "within_arm_change": "the change within one arm over time",
    "sd_of_difference": "the SD of the paired difference",
    "se_of_difference": "the SE of the paired difference",
    "n_pairs": "the number of paired measurements",
    "met": "whether the criterion is met",
    "method": "the method",
    "data": "the numbers for this arm/condition",
    "effect_estimates": "a between-group effect estimate",
    "paired_effects": "a paired comparison between two conditions",
}
SECTION = {
    "design": "Study design", "population": "Patients", "intervention": "Intervention",
    "comparator": "Comparator", "eligibility": "Eligibility for our review",
    "feasibility": "Feasibility", "citation": "Citation", "extraction_flags": "Cohort flags",
    "subgroup_flags": "Subgroup flags",
}


def describe_field(field):
    """`outcomes[mortality @28d].data[pes guided].events` -> words a clinician can read."""
    if " and " in field:  # a concern naming two entries at once; already in words
        return field
    m = re.match(r"^(outcomes|arms)\[([^\]]*)\](.*)$", field)
    if m:
        kind, label, rest = m.groups()
        when = ""
        if " @" in label:
            label, when = label.split(" @", 1)
        head = (f"Outcome '{label.strip()}'" + (f", at {when}" if when else "")) if kind == "outcomes" \
            else f"Arm/condition '{label.strip()}'"
        if not rest:
            return head + " — the whole entry"
        arm = re.search(r"\.data\[([^\]]*)\]", rest)
        paired = re.search(r"\.paired_effects\[([^\]]*)\]", rest)
        effect = re.search(r"\.effect_estimates\[([^\]]*)\]", rest)
        leaf = re.sub(r"\[[^\]]*\]", "", rest).strip(".")
        for prefix in ("data.", "paired_effects.", "effect_estimates."):
            leaf = leaf.replace(prefix, "")
        what = LEAF.get(leaf) or LEAF.get(leaf.split(".")[-1]) or leaf.replace("_", " ")
        where = ""
        if arm:
            where = f", for arm/condition '{arm.group(1)}'"
        elif paired:
            where = f", in the paired comparison '{paired.group(1)}'"
        elif effect:
            where = f", in the effect estimate '{effect.group(1)}'"
        return f"{head}: {what}{where}"
    parts = field.split(".")
    section = SECTION.get(parts[0], parts[0].replace("_", " ").capitalize())
    rest = ".".join(parts[1:])
    rest = re.sub(r"\[(\d+)\]", lambda x: f" (item {int(x.group(1)) + 1})", rest)
    rest = re.sub(r"\[([^\]]*)\]", r" '\1'", rest)
    return f"{section} → {rest.replace('_', ' ').replace('.', ' → ')}" if rest else section


def _clip(s, limit):
    return s if len(s) <= limit else s[: limit - 1] + "…"


def _n(x):
    return "—" if x is None else x


def _p(x):
    p, pv = x.get("p_value"), x.get("p_value_verbatim")
    return f", p {pv or p}" if (p is not None or pv) else ""


def _entry(d):
    """One outcome data[] entry as text."""
    arm = d.get("arm_label") or "?"
    if d.get("events") is not None or d.get("total") is not None:
        return f"{arm}: {_n(d.get('events'))}/{_n(d.get('total'))} patients"
    s = f"{arm}: {_n(d.get('value'))}"
    disp, lo, hi = d.get("dispersion"), d.get("dispersion_low"), d.get("dispersion_high")
    dtype = (d.get("dispersion_type") or "")
    if disp is not None:
        s += f" ± {disp}" if dtype.lower() in ("sd", "se", "sem") else f" ({disp})"
    elif lo is not None or hi is not None:
        s += f" ({_n(lo)} to {_n(hi)})"
    if dtype:
        s += f" {dtype}"
    if d.get("value_type") and d["value_type"] != "mean":
        s += f" [{d['value_type']}]"
    if d.get("unit"):
        s += f" {d['unit']}"
    if d.get("n") is not None:
        s += f", n={d['n']}"
    if d.get("derived"):
        s += " (computed)"
    return s


def describe_value(v, limit=900):
    """A reviewer's reading — scalar, outcome, arm, or any object — in words."""
    if v is None or v is D.MISSING:
        return "— (nothing recorded)"
    if isinstance(v, bool):
        return "yes" if v else "no"
    if isinstance(v, (int, float, str)):
        return str(v)
    if isinstance(v, list):
        return "\n\n".join(describe_value(x, limit) for x in v) if v else "— (empty list)"
    if isinstance(v, dict):
        if "name_verbatim" in v or "domain" in v:  # an outcome
            tp = v.get("timepoint") or {}
            when = (tp.get("label_verbatim")
                    or (f"{tp['hours']} h" if tp.get("hours") is not None else None)
                    or (f"day {tp['days']}" if tp.get("days") is not None else None)
                    or (tp.get("type") or "?").replace("_", " "))
            at = "" if re.match(r"^(on|at|during|after|within|before|pre|post)\b", str(when), re.I) else "at "
            lines = [f"Outcome '{_n(v.get('name_verbatim'))}' ({(v.get('domain') or '?').replace('_', ' ')}), "
                     f"{at}{when}" + ("; flagged as the paper's primary outcome" if v.get("is_primary_outcome") else "")]
            lines += ["  " + _entry(d) for d in v.get("data") or []]
            paired = v.get("paired_effects") or []
            for p in paired[:4]:
                lines.append(f"  paired: {_n(p.get('arm_label_a'))} vs {_n(p.get('arm_label_b'))}: "
                             f"difference {_n(p.get('mean_difference'))}{_p(p)}")
            if len(paired) > 4:
                lines.append(f"  … and {len(paired) - 4} more paired comparisons")
            for e in (v.get("effect_estimates") or [])[:3]:
                ci = (f" ({_n(e.get('ci_low'))} to {_n(e.get('ci_high'))})"
                      if e.get("ci_low") is not None or e.get("ci_high") is not None else "")
                lines.append(f"  effect: {_n(e.get('type'))} {_n(e.get('value'))}{ci}{_p(e)}")
            return _clip("\n".join(lines), limit)
        if "label" in v and "role" in v:  # an arm
            s = f"Arm/condition '{_n(v.get('label'))}' ({_n(v.get('role'))})"
            n = v.get("n_analysed")
            n = n.get("value") if isinstance(n, dict) else n
            if n is not None:
                s += f", n analysed {n}"
            if v.get("sequence_position") is not None:
                s += f", position {v['sequence_position']} in the order applied"
            if v.get("description_verbatim"):
                s += f" — {v['description_verbatim']}"
            return _clip(s, limit)
        if "value" in v and set(v) <= {"value", "derived", "derivation", "provenance", "value_type",
                                        "unit", "n", "total", "dispersion", "dispersion_type"}:
            s = str(_n(v.get("value")))
            if v.get("unit"):
                s += f" {v['unit']}"
            if v.get("derived"):
                s += f" (computed: {v.get('derivation')})"
            return s
        items = [f"{k.replace('_', ' ')}: {describe_value(x, 160)}" for k, x in v.items()
                 if x not in (None, [], {}) and k != "provenance" and not k.endswith("_provenance")]
        return _clip("; ".join(items), limit)
    return str(v)


def reading_text(display, value):
    if display == "<absent>":
        return "— (nothing recorded)"
    if display == "null":
        return "— (recorded as not reported)"
    return describe_value(value)


def need(row, clarifications):
    """What we need from you, for one sheet-2 row."""
    why = row["why_here"] or ""
    if why.startswith("returned Unclear"):
        for c in clarifications:
            if c["study_id"] == row["study_id"] and row["field"].startswith(c["field_startswith"]):
                return c["answer"]
        return ("You returned this Unclear last time and nobody answered you — sorry. The field is: "
                f"{describe_field(row['field'])}. Choose the reader whose entry matches the paper, "
                "or say in the notes what is still missing.")
    if why.startswith("model returned unclear"):
        note = why[len("model returned unclear:"):].strip()
        reason, _, evidence = note.rpartition(" | ")
        if not reason:
            reason, evidence = note, ""
        m = re.search(r"Co-author:\s*(.*)$", reason, re.S)
        head = ("A third reader (not GLM, not DeepSeek) checked both versions against the paper "
                "and thinks neither is fully right.")
        body = f"It asks you to: {m.group(1).strip()}" if m else f"Its note: {reason}"
        where = f"\nWhere to look: {evidence}" if evidence else ""
        return f"{head} {body}{where}\n{HOW_TO_ANSWER}"
    if why.startswith("panel conflict"):
        return ("The two readers read this differently. Compare the two quotes (and the table if "
                "needed) and choose who is right. If neither, choose 'Both wrong' and type the correct "
                "reading.")
    return why  # concerns.json rows are already written to the co-author


def category(why):
    why = why or ""
    if why.startswith("panel"):
        return 0
    if why.startswith("returned Unclear"):
        return 1
    if why.startswith("model returned"):
        return 2
    return 3


# ---------------------------------------------------------------- sheet 2

def element_versions(a, b):
    """{`outcomes[label]`: {"glm": [elements], "deepseek": [elements]}} for one study."""
    groups = {}
    for key in MP.ELEMENT_ARRAYS:
        for label, ea, eb in D.match_elements(a.get(key) or [], b.get(key) or []):
            g = groups.setdefault(f"{key}[{label}]", {"glm": [], "deepseek": []})
            if ea is not D.MISSING:
                g["glm"].append(ea)
            if eb is not D.MISSING:
                g["deepseek"].append(eb)
    return groups


def reading_rows(manifest, concerns, clarifications):
    field_rulings, _ = R.load()
    rows = []
    for sid in W.study_ids():
        m = manifest.get(sid, {})
        paper = f"{m.get('first_author', '')} {m.get('year', '')}".strip()
        open_, a, b = MP.open_rows(sid)
        versions = element_versions(a, b)
        model_unclear = {r["field"]: r for r in field_rulings
                         if r["study_id"] == sid and r["adjudicator_type"] == "model"
                         and r["decision"] == "unclear"}
        grouped = {}
        for r in open_:
            why = r["reserved"]
            if why and why.startswith("cohort relation"):
                continue  # sheet 1 settles every cohort field of a pending cohort
            if not why:
                u = model_unclear.get(r["field"])
                target = MP.element_of(r["field"])
                if not u and target in model_unclear:
                    # One question per outcome the model could not settle, not one
                    # row per difference beneath it.
                    grouped[target] = grouped.get(target, 0) + 1
                    continue
                if not u:
                    continue  # still with the model, not yet a human question
                why = f"model returned unclear: {u['note']}"
            rows.append({
                "study_id": sid, "author": paper, "kind": r["kind"], "field": r["field"],
                "why_here": why, "glm": r["glm_display"], "deepseek": r["deepseek_display"],
                "glm_quote": r["glm_quote"], "deepseek_quote": r["deepseek_quote"],
                "what this row is about": describe_field(r["field"]),
                "GLM says": reading_text(r["glm_display"], r["glm"]),
                "DeepSeek says": reading_text(r["deepseek_display"], r["deepseek"]),
            })
        for target, n in grouped.items():
            v = versions.get(target, {"glm": [], "deepseek": []})
            rows.append({
                "study_id": sid, "author": paper, "kind": "element", "field": target,
                "why_here": f"model returned unclear: {model_unclear[target]['note']}",
                "glm": f"GLM's version of this entry ({n} difference(s) beneath it)",
                "deepseek": "DeepSeek's version — see the reviewer records",
                "glm_quote": "", "deepseek_quote": "",
                "what this row is about": describe_field(target)
                + f" — whose version stands? ({n} differences beneath it)",
                "GLM says": describe_value(v["glm"]) if v["glm"] else "— (no such entry)",
                "DeepSeek says": describe_value(v["deepseek"]) if v["deepseek"] else "— (no such entry)",
            })
    for c in concerns:
        rows.append({**{k: "" for k in READING_COLUMNS + READING_MACHINE}, **c,
                     "what this row is about": describe_field(c["field"]) if c.get("field") else "",
                     "GLM says": c.get("glm", ""), "DeepSeek says": c.get("deepseek", "")})
    for r in rows:
        r["what we need from you"] = need(r, clarifications)
    rows.sort(key=lambda r: (r["study_id"], category(r["why_here"]), r["field"]))
    return rows


# ---------------------------------------------------------------- sheet 1

def cohort_rows(manifest, briefs):
    _, study_rulings = R.load()
    out = []
    for sid, row in RC.load_cohorts().items():
        if RC.settled(row):
            continue
        m, brief = manifest.get(sid, {}), briefs.get(sid, {})
        earlier = brief.get("earlier") or next(
            (f"{r['adjudicator']} chose {r['decision']} ({r['value']}) {r['note']}".strip()
             for r in reversed(study_rulings) if r["study_id"] == sid and r["topic"] == "cohort"), "")
        s = brief.get("suggestion") or {}
        out.append({
            "group": brief.get("group", ""), "study_id": sid, "author": m.get("first_author", ""),
            "paper": f"{m.get('first_author', '')} {m.get('year', '')}, {m.get('journal', '')}. {m.get('title', '')}",
            "registration": brief.get("registration", ""), "enrolment window": brief.get("window", ""),
            "centre": brief.get("centre", ""), "population": brief.get("population", ""),
            "what the paper says about shared patients": brief.get("statement", ""),
            "your earlier answer": earlier,
            "question for you": brief.get("question", row.get("note", "")),
            "our suggestion": brief.get("suggestion_text", ""),
            "cohort_id": s.get("cohort_id", ""),
            "overlapping_cohort_with": ";".join(s.get("overlapping_cohort_with", [])),
            "is_primary_report": "TRUE" if s.get("is_primary_report", True) else "FALSE",
        })
    out.sort(key=lambda r: (r["group"], r["study_id"]))
    return out


# ---------------------------------------------------------------- sheet 4

def earlier_unclear_rows(returned, adjudicator, manifest):
    """What became of each row the co-author returned Unclear in the first workbook."""
    from openpyxl import load_workbook

    if not returned or not Path(returned).exists():
        return []
    wb = load_workbook(returned)
    if "3. Data conflicts" not in wb.sheetnames:
        return []
    live = {(r["study_id"], R.strip_timepoints(r["field"]), r["kind"], str(r["glm"]), str(r["deepseek"]))
            for r in W.triaged_rows(manifest)}
    field_rulings, _ = R.load()
    later = {}
    for r in field_rulings:
        if r["adjudicator_type"] == "human" and r["decision"] != "unclear" and r["adjudicator"] != adjudicator:
            later.setdefault((r["study_id"], R.strip_timepoints(r["field"])), []).append(r)
    out = []
    for _, row in R.sheet_rows(wb["3. Data conflicts"]):
        if R.cell(row.get("DECISION")).lower() != "unclear":
            continue
        sid, field = R.cell(row["study_id"]), R.cell(row["field"])
        glm, ds, note = R.cell(row["glm"]), R.cell(row["deepseek"]), R.cell(row.get("Adjudicator notes"))
        key = (sid, R.strip_timepoints(field))
        if later.get(key):
            r = later[key][-1]
            chose = {"glm": "GLM's reading", "deepseek": "DeepSeek's reading",
                     "not_reported": "'not reported'"}.get(r["decision"], f"'{r['value']}'")
            what = f"{r['adjudicator']} settled it and chose {chose}. Reason: {r['note']}"
        elif not in_scope(field):
            what = ("Not needed after all. This field does not go into the analysis — we only settle "
                    "fields that do — so nobody has to rule on it. The final record keeps GLM's reading "
                    "there and is labelled 'not adjudicated'.")
        elif (sid, key[1], "conflict", glm, ds) not in live:
            if re.search(r"72|24|48|baseline|time|hour", note, re.I) or "timepoint" in field:
                what = ("You were right. The two readings came from different timepoints (for example "
                        "baseline versus 72 h) that our software had wrongly paired as one disagreement. "
                        "That bug is fixed — timepoints are now kept apart — and this row no longer exists.")
            else:
                what = ("This disagreement no longer exists after we fixed how the two readers' entries "
                        "are paired. Nothing to rule.")
        else:
            what = ("Back on sheet 2 — this time with an explanation of what the field means and what "
                    "we suggest.")
        out.append({"study_id": sid, "author": manifest.get(sid, {}).get("first_author", ""),
                    "what the row was about": describe_field(field), "GLM said": glm,
                    "DeepSeek said": ds, "your note": note, "what happened to it": what})
    return out


# ---------------------------------------------------------------- workbook

def sheet(wb, title, columns, rows, answers=(), machine=(), widths=None, options=None,
          prefill=()):
    ws = wb.create_sheet(title)
    headers = list(columns) + list(answers) + list(machine)
    ws.append(headers)
    for c in ws[1]:
        c.font, c.fill = W.HEADFONT, W.HEAD
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for r in rows:
        ws.append([r.get(c, "") for c in columns] + [r.get(c, "") for c in answers]
                  + [r.get(c, "") for c in machine])
    first_answer = len(columns) + 1
    first_machine = first_answer + len(answers)
    for row in ws.iter_rows(min_row=2):
        for i, c in enumerate(row, start=1):
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if i >= first_machine:
                c.font = MACHINE
            elif i >= first_answer and headers[i - 1] in prefill:
                c.fill = SUGGEST
    for i, w in enumerate(widths or [], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for i in range(first_machine, len(headers) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 12
    if options and answers:
        dv = DataValidation(type="list", formula1=options, allow_blank=True)
        ws.add_data_validation(dv)
        col = get_column_letter(first_answer + list(answers).index("DECISION"))
        dv.add(f"{col}2:{col}{max(ws.max_row, 2)}")
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    return ws


def start_here(wb, counts):
    ws = wb.active
    ws.title = "Start here"
    ws.column_dimensions["A"].width = 110
    lines = [
        ("Esophageal balloon meta-analysis — follow-up for the co-author", True),
        ("", False),
        (W.build_stamp(), False),  # row 3: rulings.py import reads this stamp
        ("", False),
        ("Sorry about the last version of this file. Its questions were written for the software, not for you.", False),
        ("Everything here is in plain words. The grey small-print columns on the right are for the software; ignore them.", False),
        ("", False),
        ("WHAT WE ARE TRYING TO AVOID", True),
        ("Some papers describe the same patients twice. If we pool both papers, we count those patients twice and", False),
        ("the result looks more certain than it is. So we group papers that share patients, and we count only one", False),
        ("paper from each group.", False),
        ("", False),
        ("THREE WORDS", True),
        ("  Cohort          one group of patients, however many papers describe them. Two papers are in the same", False),
        ("                  cohort if they share ANY patients — all of them or only some.", False),
        ("  Overlap         the same thing seen from one paper: 'this paper shares patients with that one'.", False),
        ("  Primary report  the ONE paper in a cohort whose numbers we pool. The other papers of the cohort are listed", False),
        ("                  but not pooled. A paper that shares no patients with any other is its own primary.", False),
        ("", False),
        ("HOW TO TELL WHETHER TWO PAPERS SHARE PATIENTS", True),
        ("Same hospital and overlapping dates are NOT enough on their own. In order of weight:", False),
        ("  1. The trial registration number (NCT..., DRKS...). Same number = same patients, unless a paper says the", False),
        ("     samples were different. Different numbers = different patients, unless a paper says patients were shared.", False),
        ("  2. A sentence in the paper: 'these patients were reported before', or 'none of these patients took part", False),
        ("     in any other trial'.", False),
        ("  3. Enrolment dates, centre and eligibility criteria — supporting evidence only.", False),
        ("", False),
        (f"SHEET 1 — Cohorts: {counts['cohorts']} papers in {counts['groups']} groups", True),
        ("We pulled the registration numbers, dates and the relevant sentences out of each paper for you, and filled", False),
        ("our suggested answer in the grey cells. For each row: read the evidence, change the grey cells if you disagree,", False),
        ("then set DECISION to 'Settled'. If you truly cannot tell, choose 'Still unclear' and say why in the notes.", False),
        ("  cohort_id                a short name for the group (leave blank if the paper shares no patients with any other)", False),
        ("  overlapping_cohort_with  the other papers of the group, by their id, separated by ';'", False),
        ("  is_primary_report        TRUE if this is the paper we count; FALSE if another paper of the group is", False),
        ("", False),
        (f"SHEET 2 — Readings: {counts['readings']} rows", True),
        ("Each row is one place where the two AI readers disagreed and a person must settle. Read the two columns", False),
        ("'what this row is about' and 'what we need from you' first; they say what the field means and what to do.", False),
        ("Then put your answer in DECISION:", False),
        ("  GLM / DeepSeek   that reader is right.", False),
        ("  Both wrong       neither is right: write the correct entry in 'Adjudicated value' — in words if needed,", False),
        ("                   e.g. \"DeepSeek's row, but timepoint = at protocol start\".", False),
        ("  Not reported     the paper does not report this.", False),
        ("  Unclear          you cannot tell: say in the notes what is missing, and we will answer this time.", False),
        ("Record what the paper reports. No unit conversion, no median-to-mean, no imputation.", False),
        ("", False),
        (f"SHEET 3 — Questions: {counts['questions']} rules that apply to several papers at once", True),
        ("Answer in words. Our suggestion sits next to each question.", False),
        ("", False),
        (f"SHEET 4 — Your {counts['earlier']} 'Unclear' answers from the first workbook", True),
        ("Nothing was lost. This sheet says what happened to each one. Nothing to fill in.", False),
    ]
    for text, bold in lines:
        ws.append([text])
        if bold:
            ws.cell(ws.max_row, 1).font = BOLD
    ws["A1"].font = Font(bold=True, size=12)


def load_json(path):
    return json.loads(Path(path).read_text(encoding="utf-8")) if path and Path(path).exists() else None


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=str(ROOT / "adjudication_followup.xlsx"))
    ap.add_argument("--briefs", type=Path, default=FOLLOWUP / "cohort_briefs.json")
    ap.add_argument("--clarifications", type=Path, default=FOLLOWUP / "clarifications.json")
    ap.add_argument("--concerns", type=Path, default=FOLLOWUP / "concerns.json")
    ap.add_argument("--questions", type=Path, default=FOLLOWUP / "questions.json")
    ap.add_argument("--returned", type=Path,
                    default=max(RETURNED.glob("*.xlsx"), default=None, key=lambda p: p.stat().st_mtime),
                    help="the first workbook as the co-author returned it (sheet 4)")
    ap.add_argument("--adjudicator", default="Anton Nikouline",
                    help="who returned that workbook; their own rulings are not 'later' rulings")
    args = ap.parse_args()

    manifest = W.load_manifest()
    briefs = load_json(args.briefs) or {}
    clarifications = load_json(args.clarifications) or []
    concerns = load_json(args.concerns) or []
    questions = load_json(args.questions) or []
    readings = reading_rows(manifest, concerns, clarifications)
    cohorts = cohort_rows(manifest, briefs)
    earlier = earlier_unclear_rows(args.returned, args.adjudicator, manifest)
    missing_brief = [r["study_id"] for r in cohorts if r["study_id"] not in briefs]
    if missing_brief:
        print(f"warning: no cohort brief for {', '.join(missing_brief)} — the ledger note is shown instead",
              file=sys.stderr)

    wb = Workbook()
    start_here(wb, {"cohorts": len(cohorts), "groups": len({r["group"] for r in cohorts}),
                    "readings": len(readings), "questions": len(questions), "earlier": len(earlier)})
    sheet(wb, "1. Cohort assignment", COHORT_COLUMNS, cohorts, COHORT_ANSWERS, (),
          [7, 16, 11, 46, 22, 22, 22, 30, 46, 30, 46, 46, 18, 24, 12, 14, 36],
          '"Settled,Still unclear"', prefill=("cohort_id", "overlapping_cohort_with", "is_primary_report"))
    sheet(wb, "2. Readings to rule", READING_COLUMNS, readings, READING_ANSWERS, READING_MACHINE,
          [16, 14, 40, 60, 44, 44, 40, 40, 14, 26, 34], W.DECISION)
    q = wb.create_sheet("3. Questions")
    q.append(["study_id", "question", "technical detail", "our suggestion", "Answer"])
    for c in q[1]:
        c.font, c.fill = W.HEADFONT, W.HEAD
    for item in questions:
        q.append([item.get("study_id", ""), item["question"], item.get("detail", ""),
                  item.get("suggestion", ""), ""])
    for col, w in zip("ABCDE", (18, 70, 50, 50, 50)):
        q.column_dimensions[col].width = w
    for row in q.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    q.freeze_panes = "A2"
    sheet(wb, "4. Your earlier Unclear rows", EARLIER_COLUMNS, earlier, (), (),
          [16, 12, 44, 14, 14, 36, 70])

    out = Path(args.out).expanduser()
    wb.save(out)
    print(f"wrote {out}: {len(cohorts)} cohort row(s), {len(readings)} reading(s), "
          f"{len(questions)} question(s), {len(earlier)} earlier Unclear row(s)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
